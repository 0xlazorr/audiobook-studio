"""
Universal Document Processor:
Extracts, cleans, and segments text from PDF, EPUB, DOCX, TXT, Markdown, HTML, and RTF files.
"""

import os
import re
from typing import List, Dict, Any, Optional
import pypdf
import bs4
from bs4 import BeautifulSoup


def clean_text_chunk(text: str) -> str:
    """Cleans raw extracted text for natural TTS reading."""
    if not text:
        return ""

    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Fix hyphenated line breaks (e.g., 'infor-\nmation' -> 'information')
    text = re.sub(r"(\b[A-Za-z]+)-\n([A-Za-z]+\b)", r"\1\2", text)

    # Remove lone page numbers and headers
    lines = text.split("\n")
    cleaned_lines = []
    for line in lines:
        stripped = line.strip()
        if re.match(r"^[-—–]?\s*\d+\s*[-—–]?$", stripped):
            continue
        if re.match(r"^Page\s+\d+(\s+of\s+\d+)?$", stripped, re.IGNORECASE):
            continue
        cleaned_lines.append(line)
    
    text = "\n".join(cleaned_lines)

    # Standardize quotes and dashes
    text = text.replace("“", '"').replace("”", '"')
    text = text.replace("‘", "'").replace("’", "'")
    text = text.replace("—", " — ").replace("–", " — ")

    # Collapse multiple spaces (preserve newlines)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    
    return text.strip()


def estimate_duration_minutes(word_count: int, wpm: int = 150) -> float:
    """Estimates duration in minutes at standard speaking pace."""
    if wpm <= 0:
        wpm = 150
    return round(word_count / wpm, 1)


class PDFProcessor:
    """Universal document extractor and chapter segmenter."""

    @staticmethod
    def extract_document(file_path: str, original_filename: Optional[str] = None) -> Dict[str, Any]:
        """Automatically routes to the appropriate parser based on file extension."""
        name = original_filename or file_path
        ext = os.path.splitext(name)[1].lower()
        base_title = os.path.splitext(os.path.basename(name))[0].replace("_", " ").replace("-", " ").title()

        if ext == ".pdf":
            data = PDFProcessor.extract_from_pdf(file_path)
        elif ext == ".epub":
            data = PDFProcessor.extract_from_epub(file_path)
        elif ext == ".docx":
            data = PDFProcessor.extract_from_docx(file_path)
        elif ext in [".html", ".htm"]:
            data = PDFProcessor.extract_from_html(file_path)
        elif ext == ".rtf":
            data = PDFProcessor.extract_from_rtf(file_path)
        else:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            data = PDFProcessor.extract_from_text(content, title=base_title)

        if not data.get("title") or re.match(r"^[0-9a-fA-F\-]{32,36}$", data.get("title", "")) or data.get("title") in ["Audiobook", "Document", "Untitled"]:
            data["title"] = base_title

        return data

    @staticmethod
    def extract_from_pdf(pdf_path: str) -> Dict[str, Any]:
        reader = pypdf.PdfReader(pdf_path)
        total_pages = len(reader.pages)
        
        title = "Audiobook"
        author = "Unknown Author"
        if reader.metadata:
            if reader.metadata.title:
                title = str(reader.metadata.title).strip()
            if reader.metadata.author:
                author = str(reader.metadata.author).strip()

        base_name = os.path.splitext(os.path.basename(pdf_path))[0]
        if not title or title.lower() in ["untitled", "audiobook", "document"]:
            title = base_name.replace("_", " ").replace("-", " ").title()

        pages_text = []
        for i, page in enumerate(reader.pages):
            try:
                page_raw = page.extract_text() or ""
                pages_text.append((i + 1, clean_text_chunk(page_raw)))
            except Exception as e:
                pages_text.append((i + 1, f"[Page {i+1}]"))

        # Bookmarks -> Regex -> Length
        chapters = PDFProcessor._extract_by_bookmarks(reader, pages_text)
        if not chapters or len(chapters) <= 1:
            chapters = PDFProcessor._extract_by_regex(pages_text)
        if not chapters or len(chapters) <= 1:
            chapters = PDFProcessor._extract_by_length(pages_text)

        return PDFProcessor._finalize_chapters(chapters, title, author, total_pages)

    @staticmethod
    def extract_from_epub(epub_path: str) -> Dict[str, Any]:
        """Extracts chapters cleanly from EPUB eBooks."""
        import ebooklib
        from ebooklib import epub

        book = epub.read_epub(epub_path)
        
        # Metadata
        title_meta = book.get_metadata('DC', 'title')
        title = title_meta[0][0] if title_meta else "Audiobook"
        creator_meta = book.get_metadata('DC', 'creator')
        author = creator_meta[0][0] if creator_meta else "Unknown Author"

        chapters = []
        ch_idx = 1

        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), 'html.parser')
                
                # Check for chapter title
                heading = soup.find(['h1', 'h2', 'h3'])
                h_text = heading.get_text().strip() if heading else None

                # Extract paragraph text
                paras = [p.get_text().strip() for p in soup.find_all(['p', 'div']) if p.get_text().strip()]
                raw_text = "\n\n".join(paras) if paras else soup.get_text()
                cleaned = clean_text_chunk(raw_text)

                if len(cleaned.split()) > 40:  # Ignore tiny copyright notices or blanks
                    ch_title = h_text if h_text and len(h_text) < 80 else f"Chapter {ch_idx}"
                    chapters.append({
                        "title": ch_title,
                        "content": cleaned
                    })
                    ch_idx += 1

        if not chapters:
            # Fallback to single text
            all_text = []
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    all_text.append(soup.get_text())
            full_text = "\n\n".join(all_text)
            return PDFProcessor.extract_from_text(full_text, title=title)

        return PDFProcessor._finalize_chapters(chapters, title, author, total_pages=len(chapters))

    @staticmethod
    def extract_from_docx(docx_path: str) -> Dict[str, Any]:
        """Extracts text and headings from Microsoft Word documents."""
        import docx

        doc = docx.Document(docx_path)
        base_name = os.path.splitext(os.path.basename(docx_path))[0].replace("_", " ").title()
        
        chapters = []
        current_title = "Chapter 1"
        current_paras = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if p.style and p.style.name.startswith("Heading"):
                if current_paras:
                    chapters.append({
                        "title": current_title,
                        "content": "\n\n".join(current_paras)
                    })
                    current_paras = []
                current_title = text
            else:
                current_paras.append(text)

        if current_paras:
            chapters.append({
                "title": current_title,
                "content": "\n\n".join(current_paras)
            })

        if not chapters:
            full_text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            return PDFProcessor.extract_from_text(full_text, title=base_name)

        return PDFProcessor._finalize_chapters(chapters, base_name, "Unknown Author", total_pages=len(chapters))

    @staticmethod
    def extract_from_html(html_path: str) -> Dict[str, Any]:
        """Extracts clean text and chapters from HTML documents."""
        with open(html_path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")

        # Strip scripts and styles
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()

        title_tag = soup.find("title")
        title = title_tag.get_text().strip() if title_tag else "Document"

        paras = [p.get_text().strip() for p in soup.find_all(["p", "h1", "h2", "h3", "article", "section"]) if p.get_text().strip()]
        full_text = "\n\n".join(paras)
        return PDFProcessor.extract_from_text(full_text, title=title)

    @staticmethod
    def extract_from_rtf(rtf_path: str) -> Dict[str, Any]:
        """Extracts text from Rich Text Format (RTF) files."""
        from striprtf.striprtf import rtf_to_text
        with open(rtf_path, "r", encoding="utf-8", errors="replace") as f:
            raw_text = rtf_to_text(f.read())
        base_name = os.path.splitext(os.path.basename(rtf_path))[0].replace("_", " ").title()
        return PDFProcessor.extract_from_text(raw_text, title=base_name)

    @staticmethod
    def extract_from_text(text: str, title: str = "Audiobook") -> Dict[str, Any]:
        cleaned = clean_text_chunk(text)
        pages_dummy = [(1, cleaned)]
        chapters = PDFProcessor._extract_by_regex(pages_dummy)
        if not chapters or len(chapters) <= 1:
            chapters = PDFProcessor._extract_by_length(pages_dummy)
            if not chapters:
                chapters = [{
                    "title": "Complete Text",
                    "content": cleaned
                }]

        return PDFProcessor._finalize_chapters(chapters, title, "Unknown Author", total_pages=1)

    @staticmethod
    def _finalize_chapters(chapters: List[Dict[str, Any]], title: str, author: str, total_pages: int) -> Dict[str, Any]:
        total_words = 0
        for idx, ch in enumerate(chapters, 1):
            ch["index"] = idx
            words = len(ch["content"].split())
            ch["word_count"] = words
            ch["duration_min"] = estimate_duration_minutes(words)
            total_words += words

        total_duration = estimate_duration_minutes(total_words)

        return {
            "title": title,
            "author": author,
            "total_pages": total_pages,
            "total_words": total_words,
            "total_duration_min": total_duration,
            "chapters": chapters
        }

    @staticmethod
    def _extract_by_bookmarks(reader: pypdf.PdfReader, pages_text: List[tuple]) -> List[Dict[str, Any]]:
        try:
            outline = reader.outline
            if not outline:
                return []
            
            bookmarks = []
            def parse_outline(items):
                for item in items:
                    if isinstance(item, list):
                        parse_outline(item)
                    elif hasattr(item, "title") and hasattr(item, "page"):
                        page_num = reader.get_destination_page_number(item) + 1
                        bookmarks.append((page_num, str(item.title).strip()))

            parse_outline(outline)
            if not bookmarks:
                return []

            bookmarks.sort(key=lambda x: x[0])
            chapters = []
            for i in range(len(bookmarks)):
                start_page, title = bookmarks[i]
                end_page = bookmarks[i+1][0] if i + 1 < len(bookmarks) else len(pages_text) + 1
                chapter_pages = [text for page_num, text in pages_text if start_page <= page_num < end_page]
                content = "\n\n".join(chapter_pages).strip()
                if content:
                    chapters.append({
                        "title": title,
                        "content": content
                    })

            return chapters
        except Exception:
            return []

    @staticmethod
    def _extract_by_regex(pages_text: List[tuple]) -> List[Dict[str, Any]]:
        full_text = "\n\n".join([text for _, text in pages_text])
        chapter_pattern = re.compile(
            r"(?:\n|^)"
            r"(?:#{1,3}\s*)?"
            r"((?:CHAPTER|Chapter|ACT|BOOK|PART|SECTION)\s+(?:[0-9]+|[IVXLCDM]+|[A-Za-z]+)"
            r"(?:[:.\-–—]\s*[^\n]+)?|"
            r"PROLOGUE|EPILOGUE|PREFACE|INTRODUCTION|"
            r"#{1,2}\s+[^\n]+)"
            r"(?:\n|$)",
            re.IGNORECASE
        )

        matches = list(chapter_pattern.finditer(full_text))
        if len(matches) < 2:
            return []

        chapters = []
        for i in range(len(matches)):
            match = matches[i]
            title = match.group(1).strip()
            # Clean markdown hashes or extra spaces
            title = re.sub(r"^#+\s*", "", title)
            title = re.sub(r"\s+", " ", title)

            start_pos = match.end()
            end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(full_text)
            chunk = full_text[start_pos:end_pos].strip()
            if chunk:
                chapters.append({
                    "title": title,
                    "content": chunk
                })

        return chapters

    @staticmethod
    def _extract_by_length(pages_text: List[tuple], target_words: int = 1500) -> List[Dict[str, Any]]:
        chapters = []
        current_chunk = []
        current_words = 0
        section_idx = 1
        start_page = pages_text[0][0] if pages_text else 1

        for page_num, text in pages_text:
            if not text.strip():
                continue
            words = len(text.split())
            current_chunk.append(text)
            current_words += words

            if current_words >= target_words:
                title = f"Section {section_idx} (Pages {start_page}–{page_num})" if start_page != page_num else f"Section {section_idx} (Page {start_page})"
                chapters.append({
                    "title": title,
                    "content": "\n\n".join(current_chunk).strip()
                })
                section_idx += 1
                current_chunk = []
                current_words = 0
                start_page = page_num + 1

        if current_chunk:
            last_page = pages_text[-1][0] if pages_text else start_page
            title = f"Section {section_idx} (Pages {start_page}–{last_page})" if start_page != last_page else f"Section {section_idx} (Page {start_page})"
            chapters.append({
                "title": title,
                "content": "\n\n".join(current_chunk).strip()
            })

        return chapters
